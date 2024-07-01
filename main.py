# -*- coding: utf-8 -*-
import argparse
import re
from docx import Document
import pyperclip


def find_and_replace_paragraphs(doc: Document, text2find: str, text2insert: str) -> None:
    """
    Find and replace text in all paragraphs of the document.

    Args:
    - doc (Document): The Document object from python-docx.
    - text2find (str): The text to find within each paragraph.
    - text2insert (str): The text to replace `text2find` with.
    """
    for paragraph in doc.paragraphs:
        if text2find in paragraph.text:
            paragraph.text = paragraph.text.replace(text2find, text2insert)


def extract_sections(text: str) -> list:
    """
    Extract sections from the input text using a specified pattern.

    Args:
    - text (str): The input text containing sections marked with patterns.

    Returns:
    - list: A list of tuples containing section titles and contents.
    """
    section_pattern = r"##+ (.+?)\n\n(.+?)(?=##+ |\Z)"  # Regex pattern to capture each section
    return re.findall(section_pattern, text, re.DOTALL)


def copy_prompt_to_clipboard(product_name: str, product_categories: list, company_name: str) -> None:
    """
    Copy the generated SEO prompt to clipboard.

    Args:
    - product_name (str): Name of the product.
    - product_categories (list of str): Categories or types of the product.
    """
    scenario = read_text_file('scenario.txt')
    prompt = f"Schreibe einen SEO Text zum Produkt {product_name}. " \
             f"Für das Produkt existieren die Kategorien {', '.join(product_categories)}."
    pyperclip.copy(scenario.replace('<company_name>', company_name) + '\n' + prompt)
    input('The generated prompt has been copied to your clipboard. Hit enter when you have copied and saved the LLM '
          'response to "response.txt"...')


def read_text_file(path: str) -> str:
    """
    Read content of .txt file from a path.

    Args:
    - path (str): Path to .txt file

    Returns:
    - str: Content read from the file.
    """
    with open(path, 'r', encoding='utf-8') as file:
        return file.read()


def generate_document(template_path: str, product_name: str, sections: list, output_path: str) -> None:
    """
    Generate a new document based on template and extracted sections.

    Args:
    - template_path (str): Path to the template document.
    - product_name (str): Name of the product.
    - sections (list of tuples): List of tuples containing section titles and contents.
    - output_path (str): Output path for the resulting SEO docx file.
    """
    doc = Document(template_path)
    find_and_replace_paragraphs(doc, text2find='<product_name>', text2insert=product_name)

    for idx, (_, section_text) in enumerate(sections):
        # Replace section placeholders in the document, removing paragraph breaks
        find_and_replace_paragraphs(doc, text2find=f'<section_{idx}>', text2insert=section_text.replace('\n\n', ' '))

    doc.save(output_path)
    print(f"Text added and new document saved as {output_path}")


def parse_arguments() -> argparse.Namespace:
    """
    Parse command line arguments.

    Returns:
    - argparse.Namespace: Namespace object containing parsed arguments.
    """
    parser = argparse.ArgumentParser(description='Create a new SEO text.')
    parser.add_argument('--company_name', type=str, help='Name of the company.')
    parser.add_argument('--product_name', type=str, help='Name of the product.')
    parser.add_argument('--product_categories', type=str, nargs='+',
                        help='Categories or types of the product (usually extracted from website filters).')
    parser.add_argument('--output_path', default='seo_example.docx', type=str,
                        help='Output path for the resulting SEO docx-File.')
    return parser.parse_args()


def main() -> None:
    """Main function to orchestrate the entire SEO document generation process."""
    args = parse_arguments()
    copy_prompt_to_clipboard(args.product_name, args.product_categories, args.company_name)
    llm_response = read_text_file('response.txt')
    sections = extract_sections(llm_response)
    generate_document(args.output_path, args.product_name, sections, args.output_path)


if __name__ == "__main__":
    main()
